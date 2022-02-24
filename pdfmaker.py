from fpdf import FPDF
import os
import PyPDF2
import docx
from warnings import warn

def docreader(filename):
    doc = docx.Document(filename)
    fullText = []
    for para in doc.paragraphs:
        fullText.append(para.text)
    return '\n'.join(fullText)

def textReader(filepath):
    text_file = open(filepath, "r")
    data = text_file.read()
    text_file.close()
    return data


def pdfreader(path):
    file=open(path,"rb")
    reader=PyPDF2.PdfFileReader(file)
    p=reader.getNumPages()
    r=""
    for i in range(int(p)):
        r+=(reader.getPage(i)).extractText()
    return r

def txt_maker(summary):
    with open('file.txt', 'w') as outfile:
        outfile.writelines(summary)


def create_pdf():
    pdf = FPDF() 
   
    pdf.add_page() 
     
    pdf.set_font('arial', size=15)
    
    f = open("file.txt", "r")

    for x in f:
        pdf.cell(200,10, txt=x,ln=1,align='C')
    
    pdf.output('output_pdf.pdf')
    # print(f"text_pdf is created..and it is saved in: {os.getcwd()}")